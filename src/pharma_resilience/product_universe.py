from __future__ import annotations

import csv
import re
import unicodedata
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, Sequence

from openpyxl import load_workbook


LICENSED_SHEET_NAME = "RUHSATLI ÜRÜNLER LİSTESİ"

CORE_HEADERS = (
    "SIRA NO",
    "BARKOD",
    "ÜRÜN ADI",
    "ETKİN MADDE",
    "ATC KODU",
    "RUHSAT SAHİBİ",
    "RUHSAT TARİHİ",
    "RUHSAT NUMARASI",
)


def normalize_text(value: object) -> str:
    """
    Deterministic text normalization for matching/audit only.

    Turkish dotless-i is normalized to i, diacritics are removed,
    punctuation becomes whitespace, and whitespace is collapsed.
    """
    if value is None:
        return ""

    text = str(value).strip().casefold()

    # Turkish dotless i is not removed by ordinary Unicode decomposition.
    text = text.replace("ı", "i")

    text = unicodedata.normalize("NFKD", text)

    text = "".join(
        ch
        for ch in text
        if unicodedata.category(ch) != "Mn"
    )

    text = re.sub(r"[^a-z0-9]+", " ", text)
    text = re.sub(r"\s+", " ", text).strip()

    return text


def tokenize(value: object) -> tuple[str, ...]:
    normalized = normalize_text(value)

    if not normalized:
        return ()

    return tuple(normalized.split())


@dataclass(frozen=True)
class AliasEntry:
    family_id: str
    canonical_ingredient: str
    alias: str
    alias_type: str
    notes: str
    tokens: tuple[str, ...]


@dataclass(frozen=True)
class IngredientMatch:
    family_id: str
    canonical_ingredient: str
    matched_aliases: tuple[str, ...]


def load_alias_registry(
    path: str | Path,
) -> tuple[AliasEntry, ...]:

    path = Path(path)

    required = {
        "family_id",
        "canonical_ingredient",
        "alias",
        "alias_type",
        "notes",
    }

    with path.open(
        "r",
        encoding="utf-8-sig",
        newline="",
    ) as f:
        rows = list(csv.DictReader(f))

    if not rows:
        raise ValueError("Alias registry is empty")

    actual = set(rows[0].keys())

    missing = required - actual

    if missing:
        raise ValueError(
            f"Alias registry missing columns: {sorted(missing)}"
        )

    result: list[AliasEntry] = []

    seen_family_alias: set[tuple[str, tuple[str, ...]]] = set()

    for row_number, row in enumerate(rows, 2):

        family_id = (row["family_id"] or "").strip()
        canonical = (row["canonical_ingredient"] or "").strip()
        alias = (row["alias"] or "").strip()
        alias_type = (row["alias_type"] or "").strip()
        notes = (row["notes"] or "").strip()

        if not family_id:
            raise ValueError(
                f"Missing family_id at alias row {row_number}"
            )

        if not canonical:
            raise ValueError(
                f"Missing canonical_ingredient at alias row {row_number}"
            )

        if not alias:
            raise ValueError(
                f"Missing alias at alias row {row_number}"
            )

        alias_tokens = tokenize(alias)

        if not alias_tokens:
            raise ValueError(
                f"Alias normalizes to empty at row {row_number}: {alias!r}"
            )

        duplicate_key = (
            family_id,
            alias_tokens,
        )

        if duplicate_key in seen_family_alias:
            raise ValueError(
                "Duplicate normalized family-alias: "
                f"{family_id} / {alias!r}"
            )

        seen_family_alias.add(duplicate_key)

        result.append(
            AliasEntry(
                family_id=family_id,
                canonical_ingredient=canonical,
                alias=alias,
                alias_type=alias_type,
                notes=notes,
                tokens=alias_tokens,
            )
        )

    return tuple(result)


def _contains_token_phrase(
    source_tokens: Sequence[str],
    phrase_tokens: Sequence[str],
) -> bool:
    """
    True only when phrase_tokens occur as an exact contiguous token phrase.

    This deliberately prevents substring collisions such as:
      citalopram    vs escitalopram
      loratadine   vs desloratadine
      cetirizine   vs levocetirizine
    """
    n = len(phrase_tokens)

    if n == 0 or n > len(source_tokens):
        return False

    target = tuple(phrase_tokens)

    for start in range(len(source_tokens) - n + 1):
        if tuple(source_tokens[start : start + n]) == target:
            return True

    return False


def match_seed_ingredients(
    active_ingredient_text: object,
    registry: Iterable[AliasEntry],
) -> tuple[IngredientMatch, ...]:
    """
    Match only explicit token/phrase-safe aliases.

    Multiple aliases of the same canonical ingredient collapse into one
    IngredientMatch. Multiple distinct canonical ingredients remain separate
    and can later be routed to the review queue.
    """
    source_tokens = tokenize(active_ingredient_text)

    if not source_tokens:
        return ()

    grouped: dict[
        tuple[str, str],
        set[str],
    ] = {}

    for entry in registry:

        if not _contains_token_phrase(
            source_tokens,
            entry.tokens,
        ):
            continue

        key = (
            entry.family_id,
            entry.canonical_ingredient,
        )

        grouped.setdefault(key, set()).add(entry.alias)

    result = [
        IngredientMatch(
            family_id=family_id,
            canonical_ingredient=canonical,
            matched_aliases=tuple(sorted(aliases)),
        )
        for (family_id, canonical), aliases
        in grouped.items()
    ]

    result.sort(
        key=lambda x: (
            x.family_id,
            x.canonical_ingredient,
        )
    )

    return tuple(result)


def locate_header_row(
    worksheet,
    required_headers: Sequence[str] = CORE_HEADERS,
    max_scan_rows: int = 10,
) -> int:
    required = {
        normalize_text(h)
        for h in required_headers
    }

    upper = min(
        worksheet.max_row,
        max_scan_rows,
    )

    for row_number in range(1, upper + 1):

        observed = {
            normalize_text(
                worksheet.cell(
                    row_number,
                    col_number,
                ).value
            )
            for col_number
            in range(1, worksheet.max_column + 1)
        }

        if required.issubset(observed):
            return row_number

    raise ValueError(
        "Could not locate TİTCK licensed-products header row"
    )


def map_core_columns(
    worksheet,
    header_row: int,
) -> dict[str, int]:

    normalized_to_column: dict[str, int] = {}

    for column_number in range(
        1,
        worksheet.max_column + 1,
    ):
        value = worksheet.cell(
            header_row,
            column_number,
        ).value

        key = normalize_text(value)

        if key:
            normalized_to_column[key] = column_number

    result: dict[str, int] = {}

    for header in CORE_HEADERS:

        normalized = normalize_text(header)

        if normalized not in normalized_to_column:
            raise ValueError(
                f"Missing required TİTCK column: {header}"
            )

        result[header] = normalized_to_column[
            normalized
        ]

    return result


def inspect_titck_source(
    path: str | Path,
) -> dict[str, object]:

    path = Path(path)

    if not path.exists():
        raise FileNotFoundError(path)

    workbook = load_workbook(
        path,
        read_only=True,
        data_only=True,
    )

    target_normalized = normalize_text(
        LICENSED_SHEET_NAME
    )

    worksheet = None

    for candidate in workbook.worksheets:
        if normalize_text(candidate.title) == target_normalized:
            worksheet = candidate
            break

    if worksheet is None:
        workbook.close()

        raise ValueError(
            f"Licensed-products sheet not found: "
            f"{LICENSED_SHEET_NAME}"
        )

    header_row = locate_header_row(
        worksheet
    )

    columns = map_core_columns(
        worksheet,
        header_row,
    )

    info = {
        "sheet": worksheet.title,
        "header_row": header_row,
        "data_start_row": header_row + 1,
        "max_row": worksheet.max_row,
        "max_column": worksheet.max_column,
        "core_columns": columns,
    }

    workbook.close()

    return info

@dataclass(frozen=True)
class TitckRecord:
    source_row: int
    sıra_no: object
    barcode: object
    product_name: str
    active_ingredient: str
    atc_code: str
    licence_holder: str
    licence_date: object
    licence_number: str


def iter_titck_records(
    path: str | Path,
):
    path = Path(path)

    wb = load_workbook(
        path,
        read_only=True,
        data_only=True,
    )

    target = normalize_text(LICENSED_SHEET_NAME)

    ws = next(
        (
            x for x in wb.worksheets
            if normalize_text(x.title) == target
        ),
        None,
    )

    if ws is None:
        wb.close()
        raise ValueError("Licensed-products sheet not found")

    header_row = locate_header_row(ws)
    columns = map_core_columns(ws, header_row)

    try:
        for r in range(header_row + 1, ws.max_row + 1):

            def value(name):
                return ws.cell(
                    r,
                    columns[name],
                ).value

            product_name = value("ÜRÜN ADI")

            if product_name is None:
                continue

            yield TitckRecord(
                source_row=r,
                sıra_no=value("SIRA NO"),
                barcode=value("BARKOD"),
                product_name=str(product_name).strip(),
                active_ingredient=str(
                    value("ETKİN MADDE") or ""
                ).strip(),
                atc_code=str(
                    value("ATC KODU") or ""
                ).strip(),
                licence_holder=str(
                    value("RUHSAT SAHİBİ") or ""
                ).strip(),
                licence_date=value("RUHSAT TARİHİ"),
                licence_number=str(
                    value("RUHSAT NUMARASI") or ""
                ).strip(),
            )
    finally:
        wb.close()

def iter_titck_records(
    path: str | Path,
):
    path = Path(path)

    wb = load_workbook(
        path,
        read_only=True,
        data_only=True,
    )

    target = normalize_text(LICENSED_SHEET_NAME)

    ws = next(
        (
            x for x in wb.worksheets
            if normalize_text(x.title) == target
        ),
        None,
    )

    if ws is None:
        wb.close()
        raise ValueError("Licensed-products sheet not found")

    header_row = locate_header_row(ws)
    columns = map_core_columns(ws, header_row)

    # 1-based Excel column -> 0-based tuple index
    idx = {
        name: col - 1
        for name, col in columns.items()
    }

    try:
        for source_row, row in enumerate(
            ws.iter_rows(
                min_row=header_row + 1,
                values_only=True,
            ),
            start=header_row + 1,
        ):
            product_name = row[idx["ÜRÜN ADI"]]

            if product_name is None:
                continue

            yield TitckRecord(
                source_row=source_row,
                sıra_no=row[idx["SIRA NO"]],
                barcode=row[idx["BARKOD"]],
                product_name=str(product_name).strip(),
                active_ingredient=str(
                    row[idx["ETKİN MADDE"]] or ""
                ).strip(),
                atc_code=str(
                    row[idx["ATC KODU"]] or ""
                ).strip(),
                licence_holder=str(
                    row[idx["RUHSAT SAHİBİ"]] or ""
                ).strip(),
                licence_date=row[idx["RUHSAT TARİHİ"]],
                licence_number=str(
                    row[idx["RUHSAT NUMARASI"]] or ""
                ).strip(),
            )
    finally:
        wb.close()

FAMILY_ATC_SCOPES = {
    "FAM_STATINS": ("C10AA",),
    "FAM_ARB": ("C09CA",),
    "FAM_SSRI": ("N06AB",),
    "FAM_DOAC": ("B01AF", "B01AE"),
    "FAM_DPP4": ("A10BH",),
    "FAM_SGLT2": ("A10BK",),
    "FAM_H1_ANTIHISTAMINE": ("R06AE", "R06AX"),
    "FAM_P2Y12": ("B01AC",),
    "FAM_TRIPTAN": ("N02CC",),
    "FAM_5HT3": ("A04AA",),
    "FAM_ANTIPSYCHOTIC": ("N05A",),
    "FAM_GLP1": ("A10BJ",),
    "FAM_BISPHOSPHONATE": ("M05BA",),
    "FAM_BENZO_SEIZURE": ("N05BA", "N05CD"),
}


def normalize_atc_code(value: object) -> str:
    if value is None:
        return ""

    return "".join(
        str(value).strip().upper().split()
    )


def atc_consistency_for_match(
    match: IngredientMatch,
    atc_code: object,
) -> str:
    code = normalize_atc_code(atc_code)

    if not code:
        return "MISSING_ATC"

    scopes = FAMILY_ATC_SCOPES.get(
        match.family_id
    )

    if not scopes:
        return "FAMILY_SCOPE_MISSING"

    if any(
        code.startswith(prefix)
        for prefix in scopes
    ):
        return "ATC_CONSISTENT"

    return "UNEXPECTED_ATC"


def classify_seeded_record(
    active_ingredient: object,
    atc_code: object,
    registry,
):
    matches = match_seed_ingredients(
        active_ingredient,
        registry,
    )

    if not matches:
        return "NO_SEED_MATCH", matches

    if len(matches) > 1:
        return "MULTIPLE_SEED_MATCH", matches

    status = atc_consistency_for_match(
        matches[0],
        atc_code,
    )

    return status, matches

# --- licence-level eligibility helpers ---


# --- licence-level eligibility helpers ---

def collapse_licence_suspension_states(records):
    """
    Collapse presentation-level records to licence-level suspension states.

    Expected record fields/attributes:
      - licence_number
      - suspension_status
      - family_id

    Returns:
      dict[licence_number, dict]
    """
    from collections import defaultdict

    grouped = defaultdict(list)

    for record in records:
        licence = str(
            getattr(record, "licence_number", "") or ""
        ).strip()

        if not licence:
            continue

        grouped[licence].append(record)

    collapsed = {}

    for licence, rows in grouped.items():
        states = {
            str(
                getattr(row, "suspension_status", "") or ""
            ).strip()
            for row in rows
        }

        families = {
            getattr(row, "family_id", None)
            for row in rows
        }
        families.discard(None)

        if states == {"0"}:
            state = "ALL_ACTIVE"

        elif (
            states
            and states.issubset({"1", "2", "3"})
        ):
            state = "ALL_SUSPENDED"

        elif (
            "0" in states
            and bool(states.intersection({"1", "2", "3"}))
        ):
            state = "MIXED"

        else:
            state = "UNKNOWN_STATUS"

        collapsed[licence] = {
            "licence_number": licence,
            "state": state,
            "families": tuple(sorted(families)),
            "presentation_count": len(rows),
        }

    return collapsed


def evaluate_family_eligibility(
    licence_records,
    *,
    min_products=6,
    min_active_licences=4,
):
    """
    Evaluate frozen candidate-family eligibility.

    Unit definition:
    one distinct non-suspended licence corresponds to one
    licence-level medicinal product. Multiple package /
    presentation rows under the same licence are collapsed.

    Returns one audit row per family.
    """
    from collections import defaultdict

    family_licences = defaultdict(set)

    for record in licence_records:
        family = getattr(record, "family_id", None)
        licence = str(
            getattr(record, "licence_number", "") or ""
        ).strip()
        state = getattr(record, "licence_state", None)

        if (
            not family
            or not licence
            or state != "ALL_ACTIVE"
        ):
            continue

        family_licences[family].add(licence)

    result = {}

    for family, licences in family_licences.items():
        product_count = len(licences)
        active_licence_count = len(licences)

        result[family] = {
            "eligible_products": product_count,
            "distinct_active_licences": active_licence_count,
            "passes_product_threshold":
                product_count >= min_products,
            "passes_licence_threshold":
                active_licence_count >= min_active_licences,
            "eligible": (
                product_count >= min_products
                and active_licence_count >= min_active_licences
            ),
        }

    return result
